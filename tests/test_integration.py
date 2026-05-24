# -*- coding: utf-8 -*-
import unittest
import os
import json
from unittest.mock import MagicMock, patch
from datetime import datetime
from openpyxl import load_workbook
from domain.models import NotaFiscal, CNPJ, NotaAuditada, Beneficiario, CategoriaFiscal
from infrastructure.services import EscritorLeitorNotasLocal
from application.use_cases import AuditarNotasUseCase, DeepTaxAdvisorUseCase, ExportarPlanilhaUseCase

class TestIntegrationPipeline(unittest.TestCase):
    CAMINHO_TEST_BRUTAS = "inspectir/data/test_brutas.json"
    CAMINHO_TEST_SAIDA = "inspectir/data/test_auditoria.json"
    def setUp(self):
        os.makedirs("inspectir/data", exist_ok=True)
        self.nota_mock = {"id": "NF-INTEG-01", "emitente": "CLINICA INTEGRADA", "cnpj": "12345678000199", "valor": 1500.00, "descricao": "CONSULTA CARDIOLOGISTA", "data": "10/05/2025"}
        EscritorLeitorNotasLocal.salvar_brutas([self.nota_mock], self.CAMINHO_TEST_BRUTAS)
    def tearDown(self):
        for caminho in [self.CAMINHO_TEST_BRUTAS, self.CAMINHO_TEST_SAIDA]:
            if os.path.exists(caminho): os.remove(caminho)
    def test_pipeline_de_auditoria_gravacao_e_leitura_local(self):
        adaptador_mock = MagicMock()
        nota_fiscal_dominio = NotaFiscal("NF-INTEG-01", "CLINICA INTEGRADA", CNPJ("12345678000199"), 1500.00, "CONSULTA", datetime.strptime("10/05/2025", "%d/%m/%Y"))
        adaptador_mock.analisar_em_lote.return_value = [NotaAuditada(nota_fiscal_dominio, True, CategoriaFiscal.SAUDE, Beneficiario("Titular", True), "Ok")]
        use_case = AuditarNotasUseCase(adaptador_mock)
        resultado = use_case.executar(self.CAMINHO_TEST_BRUTAS, self.CAMINHO_TEST_SAIDA)
        self.assertEqual(len(resultado), 1)

    def test_exportacao_excel_gera_planilha_com_colunas_e_cnpj_formatado(self):
        adaptador_mock = MagicMock()
        nf = NotaFiscal("NF-INTEG-01", "CLINICA INTEGRADA", CNPJ("12345678000199"), 1500.00, "CONSULTA", datetime.strptime("10/05/2025", "%d/%m/%Y"))
        adaptador_mock.analisar_em_lote.return_value = [NotaAuditada(nf, True, CategoriaFiscal.SAUDE, Beneficiario("Titular", True), "Ok")]
        AuditarNotasUseCase(adaptador_mock).executar(self.CAMINHO_TEST_BRUTAS, self.CAMINHO_TEST_SAIDA)
        caminho_xlsx = "inspectir/data/test_export.xlsx"
        try:
            ExportarPlanilhaUseCase.executar(self.CAMINHO_TEST_BRUTAS, self.CAMINHO_TEST_SAIDA, caminho_xlsx)
            self.assertTrue(os.path.exists(caminho_xlsx))
            ws = load_workbook(caminho_xlsx).active
            self.assertEqual([c.value for c in ws[1]], ["CNPJ", "Prestador", "Beneficiário", "Valor", "Categoria", "Data"])
            self.assertEqual(ws.cell(row=2, column=1).value, "12.345.678/0001-99")
        finally:
            if os.path.exists(caminho_xlsx):
                os.remove(caminho_xlsx)
